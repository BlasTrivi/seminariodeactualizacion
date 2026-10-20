"""Servicios: lógica de negocio aplicada sobre cada operación del CRUD.

Los servicios reciben los repositorios por inyección de dependencias y se
encargan de las validaciones que no pertenecen a una sola entidad:
unicidad (nombres, códigos, ISBN), integridad referencial (que exista el
género, la editorial, la moneda, el libro...), borrados bloqueados o en
cascada y los reportes que combinan varias entidades.

Sprint 2:

- ``ServicioCotizacionDolar.obtener_cotizaciones`` consulta la API del
  dólar (URL en ``.env``), descarga las cotizaciones y las registra en la
  base de datos (Ejercicio 07).
- ``ServicioCotizador`` suma la lista de precios bimonetaria (pesos y
  otra moneda seleccionada) y la exportación a CSV de los precios en
  todas las monedas con cotización (Ejercicio 08).

Sprint 3:

- ``ServicioScraping`` ejecuta el scraper de Cúspide (Ejercicio 03).
- ``ServicioComparadorPrecios`` compara el precio interno con el precio
  web y genera el archivo CSV de alertas (Ejercicio 04) y el reporte
  Excel de precios (Ejercicio 05).
- El decorador ``auditar`` registra cada operación del sistema en la
  tabla ``auditoria``; ``ServicioAuditoria`` permite consultarla
  (Ejercicio 06).
"""
from __future__ import annotations

import contextvars
import csv
import dataclasses
import datetime
import functools
import json
import os
import subprocess
import sys
import time
from dataclasses import dataclass, field
from pathlib import Path
from typing import (
    Any, Callable, ClassVar, Dict, Generic, List, Optional, Sequence, Tuple,
    TypeVar,
)

import requests
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from book_manager.database.connection import variable_entorno
from book_manager.entities.entities import (
    Auditoria, CotizacionDolar, Editorial, EntidadBase, Genero, Libro, Moneda,
    Precio, PrecioWeb, Serializable, Stock, TipoCotizacion,
)
from book_manager.repositories.repositories import (
    IRepositorio, RepositorioAuditoria, RepositorioCotizacionDolar,
    RepositorioEditorial,
    RepositorioGenero, RepositorioLibro, RepositorioMoneda,
    RepositorioPrecio, RepositorioPrecioWeb, RepositorioStock,
    RepositorioTipoCotizacion, Repositorios,
)

T = TypeVar("T", bound=EntidadBase)
F = TypeVar("F", bound=Callable[..., Any])

CODIGO_PESO = "ARS"
CODIGO_DOLAR = "USD"

#: Variables del ``.env`` con las URL de la API (Ejercicio 07).
VARIABLE_API_URL = "API_URL"
VARIABLE_API_URL_COTIZACIONES = "API_URL_COTIZACIONES"
TIEMPO_ESPERA_API = 15  # segundos

#: Carpetas del Sprint 3: código fuente y reportes generados.
CARPETA_PAQUETE = Path(__file__).resolve().parents[1]
CARPETA_SRC = CARPETA_PAQUETE.parent
CARPETA_REPORTES = CARPETA_PAQUETE / "reportes"
CARPETA_SCRAPER = CARPETA_PAQUETE / "scraper"
TIEMPO_MAXIMO_SCRAPING = 30 * 60  # segundos

#: Nombre que usa el sistema para cada "casa" de la API del dólar.
NOMBRES_CASAS_API: Dict[str, str] = {
    "oficial": "Oficial",
    "blue": "Blue",
    "bolsa": "MEP",
    "contadoconliqui": "CCL",
    "tarjeta": "Tarjeta",
    "mayorista": "Mayorista",
    "cripto": "Cripto",
}


# ---------------------------------------------------------------------------
# Sprint 3 - Ejercicio 06: auditoría mediante un decorador
# ---------------------------------------------------------------------------

#: Profundidad de llamadas auditadas en curso. Solo se registra la
#: operación más externa: si ``obtener_cotizaciones`` llama a ``crear``,
#: queda un único registro (el de la operación pedida por el usuario).
_PROFUNDIDAD_AUDITORIA: contextvars.ContextVar[int] = contextvars.ContextVar(
    "profundidad_auditoria", default=0)
LARGO_MAXIMO_DETALLE = 1500


class RegistroAuditoria:
    """Punto único donde el decorador ``auditar`` guarda los registros.

    Se configura una vez con el repositorio de auditoría (lo hace
    ``Servicios.desde_repositorios`` o ``main``). Si no está configurado
    las operaciones se ejecutan igual, sin registrar nada; si el registro
    falla, la operación del usuario no se ve afectada.
    """

    _repositorio: ClassVar[Optional[RepositorioAuditoria]] = None

    @classmethod
    def configurar(cls, repositorio: Optional[RepositorioAuditoria]) -> None:
        cls._repositorio = repositorio

    @classmethod
    def registrar(cls, accion: str, detalles: str) -> None:
        if cls._repositorio is None:
            return
        try:
            cls._repositorio.crear(Auditoria(
                accion, datetime.datetime.now(),
                detalles[:LARGO_MAXIMO_DETALLE]))
        except Exception as error:  # la auditoría nunca corta la operación
            print(f"[auditoría] no se pudo registrar '{accion}': {error}")


def describir(valor: Any) -> str:
    """Texto breve que describe un argumento o un resultado."""
    if isinstance(valor, Serializable):
        datos = {k: (v if len(v) <= 60 else v[:57] + "...")
                 for k, v in valor.a_dict().items() if v != ""}
        return f"{type(valor).__name__}({datos})"
    if isinstance(valor, tuple):
        return f"({', '.join(describir(v) for v in valor)})"
    if isinstance(valor, list):
        return f"{len(valor)} registro(s)"
    if isinstance(valor, Path):
        return str(valor)
    if dataclasses.is_dataclass(valor) and not isinstance(valor, type):
        texto = str(valor)
        return texto if not texto.startswith(type(valor).__name__ + "(") \
            else type(valor).__name__
    if isinstance(valor, datetime.datetime):
        return valor.strftime("%Y-%m-%d %H:%M:%S")
    return repr(valor)


def auditar(accion: str, entidad: Optional[str] = None) -> Callable[[F], F]:
    """Decorador que registra la operación en la tabla de auditoría.

    Registra la acción (por ejemplo ``"ALTA Libro"``), la fecha y los
    detalles: argumentos recibidos y resultado, o el error si la operación
    falló (en ese caso la excepción se vuelve a lanzar).

    Args:
        accion (str): Verbo de la operación (ALTA, BAJA, SCRAPING...).
        entidad (Optional[str]): Sobre qué se opera. Si no se indica se
            toma el atributo ``entidad_auditoria`` del servicio.
    """

    def decorador(funcion: F) -> F:
        @functools.wraps(funcion)
        def envoltura(*args: Any, **kwargs: Any) -> Any:
            objetivo = entidad
            argumentos = list(args)
            if objetivo is None and argumentos:
                objetivo = getattr(argumentos[0], "entidad_auditoria", None)
                argumentos = argumentos[1:]  # se omite ``self``
            nombre = f"{accion} {objetivo}".strip() if objetivo else accion
            partes = [describir(a) for a in argumentos]
            partes += [f"{k}={describir(v)}" for k, v in kwargs.items()]
            entrada = f"{funcion.__name__}({', '.join(partes)})"
            marca = _PROFUNDIDAD_AUDITORIA.set(
                _PROFUNDIDAD_AUDITORIA.get() + 1)
            externa = _PROFUNDIDAD_AUDITORIA.get() == 1
            try:
                resultado = funcion(*args, **kwargs)
            except Exception as error:
                if externa:
                    RegistroAuditoria.registrar(
                        f"{nombre} (ERROR)", f"{entrada} -> {error}")
                raise
            finally:
                _PROFUNDIDAD_AUDITORIA.reset(marca)
            if externa:
                RegistroAuditoria.registrar(
                    nombre, f"{entrada} -> {describir(resultado)}")
            return resultado

        return envoltura  # type: ignore[return-value]

    return decorador


# ---------------------------------------------------------------------------
# Servicio genérico
# ---------------------------------------------------------------------------

class ServicioBase(Generic[T]):
    """CRUD con validaciones. Las subclases redefinen los métodos hook."""

    nombre_entidad = "registro"
    #: Nombre con el que las operaciones quedan en la auditoría.
    entidad_auditoria = "Registro"

    def __init__(self, repositorio: IRepositorio[T]) -> None:
        self._repositorio = repositorio

    @auditar("LISTADO")
    def listar(self) -> List[T]:
        return self._repositorio.leer_todos()

    def obtener(self, id: int) -> T:
        """Devuelve la entidad o lanza ValueError si no existe."""
        entidad = self._repositorio.leer_por_id(id)
        if entidad is None:
            raise ValueError(
                f"No existe {self.nombre_entidad} con id {id}.")
        return entidad

    @auditar("ALTA")
    def crear(self, entidad: T) -> T:
        self._validar(entidad)
        return self._repositorio.crear(entidad)

    @auditar("MODIFICACION")
    def actualizar(self, entidad: T) -> T:
        if entidad.id is None:
            raise ValueError("La entidad a actualizar debe tener id.")
        self.obtener(entidad.id)
        self._validar(entidad)
        return self._repositorio.actualizar(entidad)

    @auditar("BAJA")
    def eliminar(self, id: int) -> bool:
        self._validar_eliminacion(id)
        return self._repositorio.eliminar(id)

    def _validar(self, entidad: T) -> None:
        """Hook: reglas de negocio previas a crear o actualizar."""

    def _validar_eliminacion(self, id: int) -> None:
        """Hook: reglas de negocio previas a eliminar."""

    def _validar_nombre_unico(self, entidad: T, nombre: str) -> None:
        """Comprueba que no exista otra entidad con el mismo nombre."""
        for existente in self.listar():
            mismo_nombre = getattr(existente, "nombre", "").lower()
            if mismo_nombre == nombre.lower() and existente.id != entidad.id:
                raise ValueError(f"Ya existe {self.nombre_entidad} con el "
                                 f"nombre '{nombre}'.")


# ---------------------------------------------------------------------------
# Servicios de catálogo
# ---------------------------------------------------------------------------

class ServicioGenero(ServicioBase[Genero]):
    """Reglas de negocio para géneros."""

    nombre_entidad = "un género"
    entidad_auditoria = "Genero"

    def __init__(self, repositorio: RepositorioGenero,
                 repo_libros: RepositorioLibro) -> None:
        super().__init__(repositorio)
        self._repo_libros = repo_libros

    def _validar(self, entidad: Genero) -> None:
        self._validar_nombre_unico(entidad, entidad.nombre)

    def _validar_eliminacion(self, id: int) -> None:
        cantidad = len(self._repo_libros.leer_por_genero(id))
        if cantidad:
            raise ValueError(
                f"No se puede eliminar: hay {cantidad} libro(s) del género.")


class ServicioEditorial(ServicioBase[Editorial]):
    """Reglas de negocio para editoriales."""

    nombre_entidad = "una editorial"
    entidad_auditoria = "Editorial"

    def __init__(self, repositorio: RepositorioEditorial,
                 repo_libros: RepositorioLibro) -> None:
        super().__init__(repositorio)
        self._repo_libros = repo_libros

    def _validar(self, entidad: Editorial) -> None:
        self._validar_nombre_unico(entidad, entidad.nombre)

    def _validar_eliminacion(self, id: int) -> None:
        cantidad = len(self._repo_libros.leer_por_editorial(id))
        if cantidad:
            raise ValueError(
                f"No se puede eliminar: hay {cantidad} libro(s) de la "
                "editorial.")


class ServicioMoneda(ServicioBase[Moneda]):
    """Reglas de negocio para monedas."""

    nombre_entidad = "una moneda"
    entidad_auditoria = "Moneda"

    def __init__(self, repositorio: RepositorioMoneda,
                 repo_precios: RepositorioPrecio,
                 repo_tipos: RepositorioTipoCotizacion) -> None:
        super().__init__(repositorio)
        self._repo_monedas = repositorio
        self._repo_precios = repo_precios
        self._repo_tipos = repo_tipos

    def obtener_por_codigo(self, codigo: str) -> Moneda:
        moneda = self._repo_monedas.leer_por_codigo(codigo)
        if moneda is None:
            raise ValueError(f"No existe la moneda '{codigo.upper()}'.")
        return moneda

    def _validar(self, entidad: Moneda) -> None:
        for existente in self.listar():
            if (existente.codigo == entidad.codigo
                    and existente.id != entidad.id):
                raise ValueError(
                    f"Ya existe una moneda con código {entidad.codigo}.")

    def _validar_eliminacion(self, id: int) -> None:
        cantidad = len(self._repo_precios.leer_por_moneda(id))
        if cantidad:
            raise ValueError(
                f"No se puede eliminar: hay {cantidad} precio(s) en esa "
                "moneda.")
        cantidad = len(self._repo_tipos.leer_por_moneda(id))
        if cantidad:
            raise ValueError(
                f"No se puede eliminar: hay {cantidad} tipo(s) de "
                "cotización de esa moneda.")


class ServicioTipoCotizacion(ServicioBase[TipoCotizacion]):
    """Reglas de negocio para tipos de cotización."""

    nombre_entidad = "un tipo de cotización"
    entidad_auditoria = "TipoCotizacion"

    def __init__(self, repositorio: RepositorioTipoCotizacion,
                 repo_cotizaciones: RepositorioCotizacionDolar,
                 repo_monedas: RepositorioMoneda) -> None:
        super().__init__(repositorio)
        self._repo_tipos = repositorio
        self._repo_cotizaciones = repo_cotizaciones
        self._repo_monedas = repo_monedas

    def listar_por_moneda(self, moneda_id: int) -> List[TipoCotizacion]:
        return self._repo_tipos.leer_por_moneda(moneda_id)

    def _validar(self, entidad: TipoCotizacion) -> None:
        if self._repo_monedas.leer_por_id(entidad.moneda_id) is None:
            raise ValueError(
                f"No existe una moneda con id {entidad.moneda_id}.")
        existente = self._repo_tipos.leer_por_nombre(entidad.nombre,
                                                     entidad.moneda_id)
        if existente is not None and existente.id != entidad.id:
            raise ValueError(
                f"Ya existe el tipo de cotización '{entidad.nombre}' para "
                "esa moneda.")

    def _validar_eliminacion(self, id: int) -> None:
        cantidad = len(self._repo_cotizaciones.leer_historico_por_tipo(id))
        if cantidad:
            raise ValueError(
                f"No se puede eliminar: hay {cantidad} cotización(es) "
                "registradas de ese tipo.")


@dataclass
class LibroDetalle:
    """Libro junto con los objetos relacionados (relación entre objetos)."""

    libro: Libro
    genero: Optional[Genero]
    editorial: Optional[Editorial]
    precios: List[Precio]
    stock: Optional[Stock]


class ServicioLibro(ServicioBase[Libro]):
    """Reglas de negocio para libros."""

    nombre_entidad = "un libro"
    entidad_auditoria = "Libro"

    def __init__(self, repositorio: RepositorioLibro,
                 repo_generos: RepositorioGenero,
                 repo_editoriales: RepositorioEditorial,
                 repo_precios: RepositorioPrecio,
                 repo_stock: RepositorioStock) -> None:
        super().__init__(repositorio)
        self._repo_libros = repositorio
        self._repo_generos = repo_generos
        self._repo_editoriales = repo_editoriales
        self._repo_precios = repo_precios
        self._repo_stock = repo_stock

    @auditar("BUSQUEDA")
    def buscar_por_titulo(self, texto: str) -> List[Libro]:
        return self._repo_libros.buscar_por_titulo(texto)

    @auditar("CONSULTA")
    def obtener_detalle(self, id: int) -> LibroDetalle:
        """Arma el libro con su género, editorial, precios y stock."""
        libro = self.obtener(id)
        return LibroDetalle(
            libro=libro,
            genero=self._repo_generos.leer_por_id(libro.genero_id),
            editorial=self._repo_editoriales.leer_por_id(libro.editorial_id),
            precios=self._repo_precios.leer_por_libro(id),
            stock=self._repo_stock.leer_por_libro(id),
        )

    def _validar(self, entidad: Libro) -> None:
        if self._repo_generos.leer_por_id(entidad.genero_id) is None:
            raise ValueError(
                f"No existe un género con id {entidad.genero_id}.")
        if self._repo_editoriales.leer_por_id(entidad.editorial_id) is None:
            raise ValueError(
                f"No existe una editorial con id {entidad.editorial_id}.")
        existente = self._repo_libros.leer_por_isbn(entidad.isbn)
        if existente is not None and existente.id != entidad.id:
            raise ValueError(
                f"Ya existe un libro con ISBN {entidad.isbn}: "
                f"'{existente.titulo}'.")

    @auditar("BAJA")
    def eliminar(self, id: int) -> bool:
        """Elimina el libro y, en cascada, sus precios y su stock.

        Primero se borran los registros dependientes y luego el libro,
        de modo que las claves foráneas de la base nunca queden rotas.
        """
        self.obtener(id)
        self._repo_precios.eliminar_por_libro(id)
        self._repo_stock.eliminar(id)
        return self._repositorio.eliminar(id)


# ---------------------------------------------------------------------------
# Servicios de valuación
# ---------------------------------------------------------------------------

class ServicioPrecio(ServicioBase[Precio]):
    """Reglas de negocio para precios."""

    nombre_entidad = "un precio"
    entidad_auditoria = "Precio"

    def __init__(self, repositorio: RepositorioPrecio,
                 repo_libros: RepositorioLibro,
                 repo_monedas: RepositorioMoneda) -> None:
        super().__init__(repositorio)
        self._repo_precios = repositorio
        self._repo_libros = repo_libros
        self._repo_monedas = repo_monedas

    def listar_por_libro(self, libro_id: int) -> List[Precio]:
        return self._repo_precios.leer_por_libro(libro_id)

    def _validar(self, entidad: Precio) -> None:
        if self._repo_libros.leer_por_id(entidad.libro_id) is None:
            raise ValueError(f"No existe un libro con id {entidad.libro_id}.")
        if self._repo_monedas.leer_por_id(entidad.moneda_id) is None:
            raise ValueError(
                f"No existe una moneda con id {entidad.moneda_id}.")
        existente = self._repo_precios.leer_por_libro_y_moneda(
            entidad.libro_id, entidad.moneda_id)
        if existente is not None and existente.id != entidad.id:
            raise ValueError(
                "El libro ya tiene un precio en esa moneda "
                f"(id {existente.id}); modifíquelo en lugar de crear otro.")


class ServicioStock:
    """Reglas de negocio para el stock (clave: id de libro)."""

    entidad_auditoria = "Stock"

    def __init__(self, repositorio: RepositorioStock,
                 repo_libros: RepositorioLibro) -> None:
        self._repositorio = repositorio
        self._repo_libros = repo_libros

    @auditar("LISTADO")
    def listar(self) -> List[Stock]:
        return self._repositorio.leer_todos()

    def obtener(self, libro_id: int) -> Stock:
        stock = self._repositorio.leer_por_libro(libro_id)
        if stock is None:
            raise ValueError(f"No hay stock cargado para el libro {libro_id}.")
        return stock

    @auditar("ALTA")
    def crear(self, stock: Stock) -> Stock:
        self._validar(stock)
        return self._repositorio.crear(stock)

    @auditar("MODIFICACION")
    def actualizar(self, stock: Stock) -> Stock:
        self._validar(stock)
        return self._repositorio.actualizar(stock)

    @auditar("BAJA")
    def eliminar(self, libro_id: int) -> bool:
        return self._repositorio.eliminar(libro_id)

    @auditar("INGRESO")
    def ingresar(self, libro_id: int, unidades: int) -> Stock:
        """Registra un ingreso de mercadería."""
        stock = self.obtener(libro_id)
        stock.ingresar(unidades)
        return self._repositorio.actualizar(stock)

    @auditar("EGRESO")
    def egresar(self, libro_id: int, unidades: int) -> Stock:
        """Registra un egreso (por ejemplo, una venta)."""
        stock = self.obtener(libro_id)
        stock.egresar(unidades)
        return self._repositorio.actualizar(stock)

    @auditar("REPORTE BAJO MINIMO")
    def listar_bajo_minimo(self) -> List[Tuple[Libro, Stock]]:
        """Libros cuyo stock está por debajo del mínimo definido."""
        resultado: List[Tuple[Libro, Stock]] = []
        for stock in self.listar():
            if stock.esta_bajo_minimo():
                libro = self._repo_libros.leer_por_id(stock.libro_id)
                if libro is not None:
                    resultado.append((libro, stock))
        return resultado

    def _validar(self, stock: Stock) -> None:
        if self._repo_libros.leer_por_id(stock.libro_id) is None:
            raise ValueError(f"No existe un libro con id {stock.libro_id}.")


# ---------------------------------------------------------------------------
# Cotizaciones: CRUD + consulta a la API externa (Ejercicio 07)
# ---------------------------------------------------------------------------

@dataclass
class CotizacionImportada:
    """Resultado de registrar una cotización descargada de la API."""

    moneda: Moneda
    tipo: TipoCotizacion
    cotizacion: CotizacionDolar
    nueva: bool  # True si se creó, False si se actualizó la del día


class ClienteCotizacionesAPI:
    """Descarga las cotizaciones publicadas por la API (dolarapi.com).

    Las URL se leen del archivo ``.env`` (``API_URL`` para los distintos
    tipos de dólar y, opcionalmente, ``API_URL_COTIZACIONES`` para las
    demás monedas). Cada registro devuelto tiene las claves ``moneda``,
    ``casa``, ``nombre``, ``compra``, ``venta`` y ``fechaActualizacion``.
    """

    def __init__(self, url: Optional[str] = None,
                 url_cotizaciones: Optional[str] = None) -> None:
        self._url = url or variable_entorno(VARIABLE_API_URL)
        self._url_cotizaciones = url_cotizaciones or variable_entorno(
            VARIABLE_API_URL_COTIZACIONES, "")

    @property
    def url(self) -> str:
        return self._url

    @staticmethod
    def _descargar(url: str) -> List[Dict[str, Any]]:
        try:
            respuesta = requests.get(url, timeout=TIEMPO_ESPERA_API)
            respuesta.raise_for_status()
            datos = respuesta.json()
        except requests.RequestException as error:
            raise ValueError(
                f"No se pudo consultar la API ({url}): {error}") from error
        except ValueError as error:
            raise ValueError(
                f"La API ({url}) no devolvió un JSON válido.") from error
        if isinstance(datos, dict):
            datos = [datos]
        if not isinstance(datos, list):
            raise ValueError(f"Respuesta inesperada de la API ({url}).")
        return datos

    def buscar_dolares(self) -> List[Dict[str, Any]]:
        """Cotizaciones del dólar por casa (API_URL)."""
        return self._descargar(self._url)

    def buscar_otras_monedas(self) -> List[Dict[str, Any]]:
        """Cotizaciones de otras monedas (vacío si no está configurada)."""
        if not self._url_cotizaciones:
            return []
        registros = self._descargar(self._url_cotizaciones)
        return [registro for registro in registros
                if str(registro.get("moneda", "")).upper() != CODIGO_DOLAR]


class ServicioCotizacionDolar:
    """Reglas de negocio para cotizaciones (clave: tipo y fecha)."""

    entidad_auditoria = "Cotizacion"

    def __init__(self, repositorio: RepositorioCotizacionDolar,
                 repo_tipos: RepositorioTipoCotizacion,
                 repo_monedas: RepositorioMoneda,
                 cliente_api: Optional[ClienteCotizacionesAPI] = None
                 ) -> None:
        self._repositorio = repositorio
        self._repo_tipos = repo_tipos
        self._repo_monedas = repo_monedas
        self._cliente_api = cliente_api

    # -- CRUD -------------------------------------------------------------

    @auditar("LISTADO")
    def listar(self) -> List[CotizacionDolar]:
        return self._repositorio.leer_todos()

    @auditar("HISTORICO")
    def listar_por_tipo(self, tipo_id: int) -> List[CotizacionDolar]:
        return self._repositorio.leer_historico_por_tipo(tipo_id)

    def obtener(self, tipo_id: int,
                fecha: datetime.date) -> CotizacionDolar:
        cotizacion = self._repositorio.leer_por_tipo_y_fecha(tipo_id, fecha)
        if cotizacion is None:
            raise ValueError(
                f"No hay cotización del tipo {tipo_id} para {fecha}.")
        return cotizacion

    def obtener_ultima(self, tipo_id: int) -> CotizacionDolar:
        """Cotización más reciente de un tipo; falla si no hay ninguna."""
        if self._repo_tipos.leer_por_id(tipo_id) is None:
            raise ValueError(
                f"No existe un tipo de cotización con id {tipo_id}.")
        cotizacion = self._repositorio.leer_ultima_por_tipo(tipo_id)
        if cotizacion is None:
            raise ValueError(
                f"No hay cotizaciones cargadas para el tipo {tipo_id}.")
        return cotizacion

    @auditar("ALTA")
    def crear(self, cotizacion: CotizacionDolar) -> CotizacionDolar:
        self._validar(cotizacion)
        return self._repositorio.crear(cotizacion)

    @auditar("MODIFICACION")
    def actualizar(self, cotizacion: CotizacionDolar) -> CotizacionDolar:
        self._validar(cotizacion)
        return self._repositorio.actualizar(cotizacion)

    @auditar("BAJA")
    def eliminar(self, tipo_id: int, fecha: datetime.date) -> bool:
        return self._repositorio.eliminar(tipo_id, fecha)

    def _validar(self, cotizacion: CotizacionDolar) -> None:
        if self._repo_tipos.leer_por_id(cotizacion.tipo_id) is None:
            raise ValueError(
                f"No existe un tipo de cotización con id "
                f"{cotizacion.tipo_id}.")
        if cotizacion.fecha > datetime.date.today():
            raise ValueError("La fecha de la cotización no puede ser futura.")
        if cotizacion.compra > cotizacion.venta:
            raise ValueError(
                "El valor de compra no puede superar al de venta.")

    # -- API externa (Ejercicio 07) ---------------------------------------

    @property
    def cliente_api(self) -> ClienteCotizacionesAPI:
        if self._cliente_api is None:
            self._cliente_api = ClienteCotizacionesAPI()
        return self._cliente_api

    @auditar("API COTIZACIONES")
    def obtener_cotizaciones(self) -> List[CotizacionImportada]:
        """Busca, descarga y registra en la base las cotizaciones de la API.

        Para cada registro devuelto por la API se resuelve la moneda (se
        crea si no existe), el tipo de cotización (idem) y se guarda la
        cotización del día: si ya había una para ese tipo y fecha se
        actualiza con los valores nuevos.

        Returns:
            List[CotizacionImportada]: Una entrada por cotización
                registrada, indicando si fue alta o modificación.

        Raises:
            ValueError: Si la API no responde o devuelve datos inválidos.
        """
        registros = self.cliente_api.buscar_dolares()
        registros.extend(self.cliente_api.buscar_otras_monedas())
        importadas: List[CotizacionImportada] = []
        for registro in registros:
            importada = self._registrar_desde_api(registro)
            if importada is not None:
                importadas.append(importada)
        if not importadas:
            raise ValueError("La API no devolvió cotizaciones válidas.")
        return importadas

    def _registrar_desde_api(
            self, registro: Dict[str, Any]) -> Optional[CotizacionImportada]:
        """Convierte un registro de la API en una cotización persistida."""
        venta = registro.get("venta")
        if venta is None:
            return None  # sin precio de venta no hay cotización utilizable
        compra = registro.get("compra")
        if compra is None or float(compra) > float(venta):
            compra = venta
        moneda = self._moneda_desde_api(registro)
        tipo = self._tipo_desde_api(registro, moneda)
        fecha = self._fecha_desde_api(registro.get("fechaActualizacion"))
        cotizacion = CotizacionDolar(tipo.id or 0, fecha, float(compra),
                                     float(venta))
        existente = self._repositorio.leer_por_tipo_y_fecha(
            cotizacion.tipo_id, fecha)
        if existente is None:
            self.crear(cotizacion)
        else:
            self.actualizar(cotizacion)
        return CotizacionImportada(moneda, tipo, cotizacion,
                                   nueva=existente is None)

    def _moneda_desde_api(self, registro: Dict[str, Any]) -> Moneda:
        codigo = str(registro.get("moneda") or CODIGO_DOLAR).upper()
        moneda = self._repo_monedas.leer_por_codigo(codigo)
        if moneda is None:
            nombre = str(registro.get("nombre") or codigo)
            moneda = self._repo_monedas.crear(Moneda(codigo, nombre, ""))
        return moneda

    def _tipo_desde_api(self, registro: Dict[str, Any],
                        moneda: Moneda) -> TipoCotizacion:
        casa = str(registro.get("casa") or "").strip().lower()
        if moneda.codigo == CODIGO_DOLAR:
            nombre = NOMBRES_CASAS_API.get(
                casa, str(registro.get("nombre") or casa or "API"))
        else:
            nombre = casa.capitalize() or "Oficial"
        tipo = self._repo_tipos.leer_por_nombre(nombre, moneda.id)
        if tipo is None:
            descripcion = f"{registro.get('nombre', nombre)} (importado " \
                          f"desde la API, casa '{casa}')"
            tipo = self._repo_tipos.crear(
                TipoCotizacion(nombre, descripcion, moneda.id or 0))
        return tipo

    @staticmethod
    def _fecha_desde_api(valor: Optional[str]) -> datetime.date:
        """Convierte ``2026-10-08T18:00:00.000Z`` en una fecha local."""
        if not valor:
            return datetime.date.today()
        texto = str(valor).replace("Z", "+00:00")
        try:
            instante = datetime.datetime.fromisoformat(texto)
        except ValueError:
            return datetime.date.today()
        if instante.tzinfo is not None:
            instante = instante.astimezone()
        fecha = instante.date()
        return min(fecha, datetime.date.today())


# ---------------------------------------------------------------------------
# Reportes: valuación del catálogo según la cotización elegida
# ---------------------------------------------------------------------------

@dataclass
class LibroCotizado:
    """Resultado de cotizar un libro con un tipo de dólar."""

    libro: Libro
    precio: Precio
    moneda: Moneda
    cotizacion: CotizacionDolar
    valor_ars: float
    valor_usd: float


@dataclass
class PrecioBimonetario:
    """Precio de un libro en pesos y en otra moneda seleccionada."""

    libro: Libro
    precio: Precio
    moneda_lista: Moneda
    valor_ars: float
    moneda: Moneda
    tipo: TipoCotizacion
    cotizacion: CotizacionDolar
    valor_moneda: float


@dataclass
class ColumnaMoneda:
    """Una moneda/tipo con cotización vigente, usada en la exportación."""

    moneda: Moneda
    tipo: TipoCotizacion
    cotizacion: CotizacionDolar

    @property
    def encabezado(self) -> str:
        return f"{self.moneda.codigo}_{self.tipo.nombre}".replace(" ", "_")


class ServicioCotizador:
    """Convierte los precios de los libros entre pesos y otras monedas."""

    entidad_auditoria = "Precios"

    def __init__(self, libros: ServicioLibro, precios: ServicioPrecio,
                 monedas: ServicioMoneda, stock: ServicioStock,
                 tipos: ServicioTipoCotizacion,
                 cotizaciones: ServicioCotizacionDolar) -> None:
        self._libros = libros
        self._precios = precios
        self._monedas = monedas
        self._stock = stock
        self._tipos = tipos
        self._cotizaciones = cotizaciones

    # -- reportes del Sprint 1 --------------------------------------------

    @auditar("REPORTE COTIZACION LIBRO")
    def cotizar_libro(self, libro_id: int, tipo_id: int) -> LibroCotizado:
        """Valúa un libro en ARS y USD usando la última cotización del tipo.

        Se toma el precio en USD si existe; en caso contrario el precio en
        ARS. Cualquier otra moneda de lista no puede cotizarse.
        """
        libro = self._libros.obtener(libro_id)
        cotizacion = self._cotizacion_dolar(tipo_id)
        return self._cotizar(libro, cotizacion)

    @auditar("REPORTE CATALOGO VALUADO")
    def catalogo_valuado(self, tipo_id: int) -> List[LibroCotizado]:
        """Cotiza todos los libros que tengan precio en ARS o USD."""
        cotizacion = self._cotizacion_dolar(tipo_id)
        resultado: List[LibroCotizado] = []
        for libro in self._libros.listar():
            try:
                resultado.append(self._cotizar(libro, cotizacion))
            except ValueError:
                continue  # libro sin precio cotizable: se omite del listado
        return resultado

    @auditar("REPORTE VALOR INVENTARIO")
    def valor_inventario(self, tipo_id: int) -> Tuple[float, float]:
        """Suma precio x cantidad en stock para todo el catálogo."""
        total_ars = 0.0
        total_usd = 0.0
        for cotizado in self.catalogo_valuado(tipo_id):
            try:
                cantidad = self._stock.obtener(cotizado.libro.id or 0).cantidad
            except ValueError:
                cantidad = 0
            total_ars += cotizado.valor_ars * cantidad
            total_usd += cotizado.valor_usd * cantidad
        return round(total_ars, 2), round(total_usd, 2)

    # -- reportes del Sprint 2 (Ejercicio 08) -----------------------------

    @auditar("CONSULTA MONEDAS COTIZADAS")
    def monedas_con_cotizacion(self) -> List[ColumnaMoneda]:
        """Monedas (distintas del peso) que tienen al menos una cotización.

        Devuelve una entrada por cada tipo de cotización con datos, con
        la cotización más reciente de ese tipo.
        """
        columnas: List[ColumnaMoneda] = []
        for moneda in self._monedas.listar():
            if moneda.codigo == CODIGO_PESO:
                continue
            for tipo in self._tipos.listar_por_moneda(moneda.id or 0):
                try:
                    cotizacion = self._cotizaciones.obtener_ultima(
                        tipo.id or 0)
                except ValueError:
                    continue
                columnas.append(ColumnaMoneda(moneda, tipo, cotizacion))
        return columnas

    @auditar("REPORTE LISTA BIMONETARIA")
    def lista_precios_bimonetaria(self, tipo_dolar_id: int,
                                  tipo_moneda_id: int
                                  ) -> List[PrecioBimonetario]:
        """Lista de precios en pesos y en la moneda del tipo seleccionado.

        Args:
            tipo_dolar_id (int): Tipo de dólar con el que se pasan a pesos
                los libros cuyo precio de lista está en USD.
            tipo_moneda_id (int): Tipo de cotización (de cualquier moneda)
                en el que se expresa la segunda columna.

        Returns:
            List[PrecioBimonetario]: Un registro por libro cotizable.
        """
        tipo = self._tipos.obtener(tipo_moneda_id)
        moneda = self._monedas.obtener(tipo.moneda_id)
        cotizacion = self._cotizaciones.obtener_ultima(tipo_moneda_id)
        resultado: List[PrecioBimonetario] = []
        for cotizado in self.catalogo_valuado(tipo_dolar_id):
            resultado.append(PrecioBimonetario(
                libro=cotizado.libro, precio=cotizado.precio,
                moneda_lista=cotizado.moneda, valor_ars=cotizado.valor_ars,
                moneda=moneda, tipo=tipo, cotizacion=cotizacion,
                valor_moneda=round(cotizado.valor_ars / cotizacion.venta, 2)))
        return resultado

    @auditar("EXPORTACION CSV")
    def exportar_precios_csv(self, ruta: Path,
                             tipo_dolar_id: int) -> Tuple[Path, int, int]:
        """Exporta los precios en pesos y en todas las monedas con cotización.

        Args:
            ruta (Path): Archivo CSV de destino (se crea la carpeta).
            tipo_dolar_id (int): Tipo de dólar para pasar a pesos los
                libros con precio de lista en USD.

        Returns:
            Tuple[Path, int, int]: Ruta escrita, cantidad de libros y
                cantidad de columnas de moneda exportadas.
        """
        columnas = self.monedas_con_cotizacion()
        if not columnas:
            raise ValueError(
                "No hay cotizaciones registradas: obtenga las cotizaciones "
                "por API o cargue alguna manualmente.")
        catalogo = self.catalogo_valuado(tipo_dolar_id)
        encabezados = ["id", "isbn", "titulo", "autor", "moneda_lista",
                       "precio_lista", CODIGO_PESO]
        encabezados.extend(columna.encabezado for columna in columnas)
        ruta = Path(ruta)
        ruta.parent.mkdir(parents=True, exist_ok=True)
        with ruta.open("w", newline="", encoding="utf-8") as archivo:
            escritor = csv.writer(archivo, lineterminator="\n")
            escritor.writerow(encabezados)
            escritor.writerow(self._fila_cotizaciones(columnas))
            for cotizado in catalogo:
                fila: List[object] = [
                    cotizado.libro.id, cotizado.libro.isbn,
                    cotizado.libro.titulo, cotizado.libro.autor,
                    cotizado.moneda.codigo, f"{cotizado.precio.valor:.2f}",
                    f"{cotizado.valor_ars:.2f}"]
                fila.extend(
                    f"{cotizado.valor_ars / columna.cotizacion.venta:.2f}"
                    for columna in columnas)
                escritor.writerow(fila)
        return ruta, len(catalogo), len(columnas)

    @staticmethod
    def _fila_cotizaciones(columnas: Sequence[ColumnaMoneda]) -> List[object]:
        """Fila informativa con la cotización (venta) usada en cada columna."""
        fila: List[object] = ["", "", "Cotización de venta utilizada", "",
                              "", "", "1.00"]
        fila.extend(f"{columna.cotizacion.venta:.2f}" for columna in columnas)
        return fila

    # -- internos ---------------------------------------------------------

    def _cotizacion_dolar(self, tipo_id: int) -> CotizacionDolar:
        """Última cotización de un tipo que debe pertenecer al dólar."""
        tipo = self._tipos.obtener(tipo_id)
        moneda = self._monedas.obtener(tipo.moneda_id)
        if moneda.codigo != CODIGO_DOLAR:
            raise ValueError(
                f"El tipo '{tipo.nombre}' corresponde a {moneda.codigo}; "
                "para valuar en pesos se necesita un tipo de dólar.")
        return self._cotizaciones.obtener_ultima(tipo_id)

    def _cotizar(self, libro: Libro,
                 cotizacion: CotizacionDolar) -> LibroCotizado:
        precio, moneda = self._precio_base(libro.id or 0)
        if moneda.codigo == CODIGO_DOLAR:
            valor_usd = precio.valor
            valor_ars = round(precio.valor * cotizacion.venta, 2)
        else:
            valor_ars = precio.valor
            valor_usd = round(precio.valor / cotizacion.venta, 2)
        return LibroCotizado(libro, precio, moneda, cotizacion,
                             valor_ars, valor_usd)

    def _precio_base(self, libro_id: int) -> Tuple[Precio, Moneda]:
        """Precio en USD si existe, si no en ARS; si no hay, ValueError."""
        precios = self._precios.listar_por_libro(libro_id)
        if not precios:
            raise ValueError(f"El libro {libro_id} no tiene precio cargado.")
        candidatos: Dict[str, Tuple[Precio, Moneda]] = {}
        for precio in precios:
            moneda = self._monedas.obtener(precio.moneda_id)
            candidatos[moneda.codigo] = (precio, moneda)
        for codigo in (CODIGO_DOLAR, CODIGO_PESO):
            if codigo in candidatos:
                return candidatos[codigo]
        raise ValueError(
            f"El libro {libro_id} solo tiene precios en monedas que no se "
            "pueden cotizar (se admiten ARS y USD).")


# ---------------------------------------------------------------------------
# Sprint 3 - Ejercicio 03: scraping de precios de la competencia
# ---------------------------------------------------------------------------

@dataclass
class ResultadoScraping:
    """Resumen de una ejecución del scraper."""

    fecha_extraccion: datetime.datetime
    libros_buscados: int
    libros_con_resultados: int
    libros_sin_resultados: List[int]
    resultados_guardados: int
    libros_con_coincidencia: int
    archivo_resultados: Path
    segundos: float

    def __str__(self) -> str:
        return (f"Scraping del {self.fecha_extraccion}: "
                f"{self.libros_buscados} libro(s) buscados, "
                f"{self.resultados_guardados} resultado(s) guardados, "
                f"{self.libros_con_coincidencia} libro(s) encontrados "
                f"({self.segundos:.0f} s).")


class ServicioScraping(ServicioBase[PrecioWeb]):
    """CRUD de los precios web y ejecución del ``CuspideSpider``."""

    nombre_entidad = "un precio web"
    entidad_auditoria = "PrecioWeb"

    def __init__(self, repositorio: RepositorioPrecioWeb,
                 repo_libros: RepositorioLibro,
                 carpeta_salida: Path = CARPETA_REPORTES / "scraping"
                 ) -> None:
        super().__init__(repositorio)
        self._repo_web = repositorio
        self._repo_libros = repo_libros
        self._carpeta_salida = Path(carpeta_salida)

    def _validar(self, entidad: PrecioWeb) -> None:
        if self._repo_libros.leer_por_id(entidad.libro_id) is None:
            raise ValueError(f"No existe un libro con id {entidad.libro_id}.")

    # -- consultas ----------------------------------------------------------

    def ultima_fecha_extraccion(self) -> datetime.datetime:
        """Fecha de la última ejecución; falla si nunca se ejecutó."""
        fecha = self._repo_web.ultima_fecha_extraccion()
        if fecha is None:
            raise ValueError("Todavía no se ejecutó el scraping: use la "
                             "opción 'Ejecutar scraping'.")
        return fecha

    def resultados(self, fecha: Optional[datetime.datetime] = None
                   ) -> List[PrecioWeb]:
        """Todos los resultados de una ejecución (por defecto la última)."""
        return self._repo_web.leer_por_fecha(
            fecha or self.ultima_fecha_extraccion())

    def seleccionados(self, fecha: Optional[datetime.datetime] = None
                      ) -> List[PrecioWeb]:
        """Resultado elegido para cada libro en una ejecución."""
        return self._repo_web.leer_seleccionados(
            fecha or self.ultima_fecha_extraccion())

    # -- ejecución del spider -----------------------------------------------

    @auditar("SCRAPING")
    def ejecutar_scraping(self, libros_ids: Optional[Sequence[int]] = None
                          ) -> ResultadoScraping:
        """Ejecuta el ``CuspideSpider`` y devuelve el resumen.

        El spider corre en un proceso separado (ver ``scraper/runner.py``)
        para poder ejecutarlo varias veces en la misma sesión. Los
        resultados quedan en la tabla ``precios_web`` y en un JSON dentro
        de ``reportes/scraping``.

        Args:
            libros_ids (Optional[Sequence[int]]): Libros a buscar; por
                defecto, todo el catálogo.

        Raises:
            ValueError: Si el catálogo está vacío o el scraper falla.
        """
        if not self._repo_libros.leer_todos():
            raise ValueError("No hay libros en el catálogo para buscar.")
        fecha = datetime.datetime.now().replace(microsecond=0)
        sello = fecha.strftime("%Y%m%d_%H%M%S")
        self._carpeta_salida.mkdir(parents=True, exist_ok=True)
        archivo_resumen = self._carpeta_salida / f"resumen_{sello}.json"
        archivo_resultados = self._carpeta_salida / f"cuspide_{sello}.json"
        comando = [sys.executable, "-m", "book_manager.scraper.runner",
                   "--resumen", str(archivo_resumen),
                   "--resultados", str(archivo_resultados),
                   "--fecha", fecha.isoformat(sep=" ")]
        if libros_ids:
            comando += ["--libros", ",".join(str(i) for i in libros_ids)]
        entorno = dict(os.environ)
        entorno["PYTHONPATH"] = os.pathsep.join(
            filter(None, [str(CARPETA_SRC), entorno.get("PYTHONPATH")]))
        inicio = time.monotonic()
        proceso = subprocess.run(
            comando, cwd=CARPETA_SCRAPER, env=entorno, capture_output=True,
            text=True, timeout=TIEMPO_MAXIMO_SCRAPING, check=False)
        if proceso.returncode != 0 or not archivo_resumen.exists():
            detalle = (proceso.stderr or proceso.stdout).strip()[-1500:]
            raise ValueError(f"El scraper terminó con error:\n{detalle}")
        datos = json.loads(archivo_resumen.read_text(encoding="utf-8"))
        return ResultadoScraping(
            fecha_extraccion=fecha,
            libros_buscados=int(datos["libros_buscados"]),
            libros_con_resultados=int(datos["libros_con_resultados"]),
            libros_sin_resultados=list(datos["libros_sin_resultados"]),
            resultados_guardados=int(datos["resultados_guardados"]),
            libros_con_coincidencia=int(datos["libros_con_coincidencia"]),
            archivo_resultados=archivo_resultados,
            segundos=round(time.monotonic() - inicio, 1))


# ---------------------------------------------------------------------------
# Sprint 3 - Ejercicio 04: comparación de precios y alertas
# ---------------------------------------------------------------------------

ALERTA_COMPETIDOR_MAS_BARATO = "Competidor más barato"
ALERTA_COMPETIDOR_MAS_CARO = "Competidor más caro"


@dataclass
class ComparacionPrecio:
    """Precio interno de un libro contra el precio publicado en la web."""

    libro: Libro
    precio_interno: float
    precio_web: PrecioWeb

    @property
    def diferencia(self) -> float:
        """Precio web menos precio interno (negativo: la web es más barata)."""
        return round(self.precio_web.precio - self.precio_interno, 2)

    @property
    def diferencia_porcentual(self) -> float:
        """Diferencia expresada como porcentaje del precio interno."""
        return round(self.diferencia / self.precio_interno * 100, 2)

    @property
    def fecha_extraccion(self) -> datetime.datetime:
        return self.precio_web.fecha_extraccion

    @property
    def tipo_alerta(self) -> str:
        return (ALERTA_COMPETIDOR_MAS_BARATO if self.diferencia < 0
                else ALERTA_COMPETIDOR_MAS_CARO)

    def supera(self, umbral_porcentual: float) -> bool:
        """True si la diferencia (en valor absoluto) alcanza el umbral."""
        return abs(self.diferencia_porcentual) >= umbral_porcentual


@dataclass
class ResultadoAlertas:
    """Archivo de alertas generado y comparaciones que lo originaron."""

    archivo: Path
    umbral_porcentual: float
    comparaciones: List[ComparacionPrecio]
    alertas: List[ComparacionPrecio] = field(default_factory=list)

    def __str__(self) -> str:
        return (f"{len(self.alertas)} alerta(s) de {len(self.comparaciones)}"
                f" libro(s) comparados -> {self.archivo}")


class ServicioComparadorPrecios:
    """Compara nuestros precios con los de la competencia."""

    entidad_auditoria = "Precios web"

    ENCABEZADOS_ALERTAS = (
        "fecha_extraccion", "libro_id", "isbn", "libro", "precio_interno",
        "precio_web", "diferencia", "diferencia_porcentual",
        "umbral_porcentual", "alerta", "titulo_web", "url")

    def __init__(self, libros: ServicioLibro, monedas: ServicioMoneda,
                 tipos: ServicioTipoCotizacion, cotizador: ServicioCotizador,
                 scraping: ServicioScraping,
                 carpeta_alertas: Path = CARPETA_REPORTES / "alertas"
                 ) -> None:
        self._libros = libros
        self._monedas = monedas
        self._tipos = tipos
        self._cotizador = cotizador
        self._scraping = scraping
        self._carpeta_alertas = Path(carpeta_alertas)

    def tipo_dolar_por_defecto(self) -> int:
        """Id del dólar Oficial (o del primer tipo de dólar disponible).

        Se usa para pasar a pesos los libros cuyo precio de lista está en
        dólares, ya que Cúspide publica sus precios en pesos.
        """
        dolar = self._monedas.obtener_por_codigo(CODIGO_DOLAR)
        tipos = self._tipos.listar_por_moneda(dolar.id or 0)
        if not tipos:
            raise ValueError("No hay tipos de cotización del dólar cargados.")
        for tipo in tipos:
            if tipo.nombre.lower() == "oficial":
                return tipo.id or 0
        return tipos[0].id or 0

    def precio_interno(self, libro_id: int,
                       tipo_dolar_id: Optional[int] = None) -> float:
        """Precio interno del libro expresado en pesos."""
        tipo = tipo_dolar_id or self.tipo_dolar_por_defecto()
        return self._cotizador.cotizar_libro(libro_id, tipo).valor_ars

    @auditar("COMPARACION")
    def comparar(self, tipo_dolar_id: Optional[int] = None,
                 fecha: Optional[datetime.datetime] = None
                 ) -> List[ComparacionPrecio]:
        """Compara cada libro encontrado en la web con su precio interno.

        Args:
            tipo_dolar_id (Optional[int]): Dólar para pasar a pesos los
                precios de lista en USD (por defecto el Oficial).
            fecha (Optional[datetime.datetime]): Ejecución del scraper a
                comparar (por defecto la última).
        """
        tipo = tipo_dolar_id or self.tipo_dolar_por_defecto()
        comparaciones: List[ComparacionPrecio] = []
        for precio_web in self._scraping.seleccionados(fecha):
            try:
                libro = self._libros.obtener(precio_web.libro_id)
                interno = self.precio_interno(precio_web.libro_id, tipo)
            except ValueError:
                continue  # libro borrado o sin precio interno cotizable
            comparaciones.append(ComparacionPrecio(libro, interno,
                                                   precio_web))
        return sorted(comparaciones, key=lambda c: c.libro.id or 0)

    @auditar("ALERTAS CSV")
    def generar_alertas_csv(self, umbral_porcentual: float,
                            tipo_dolar_id: Optional[int] = None,
                            fecha: Optional[datetime.datetime] = None
                            ) -> ResultadoAlertas:
        """Genera el archivo CSV con las alertas de precios.

        Hay alerta cuando la diferencia entre el precio web y el interno
        (en valor absoluto y en porcentaje del precio interno) es mayor o
        igual al umbral indicado por el usuario.

        Args:
            umbral_porcentual (float): Diferencia de precios (%) a partir
                de la cual se genera una alerta.
            tipo_dolar_id (Optional[int]): Dólar para los precios en USD.
            fecha (Optional[datetime.datetime]): Ejecución del scraper.

        Returns:
            ResultadoAlertas: Ruta del CSV, comparaciones y alertas.
        """
        if umbral_porcentual < 0:
            raise ValueError("La diferencia de precios no puede ser negativa.")
        comparaciones = self.comparar(tipo_dolar_id, fecha)
        alertas = [c for c in comparaciones if c.supera(umbral_porcentual)]
        momento = fecha or self._scraping.ultima_fecha_extraccion()
        self._carpeta_alertas.mkdir(parents=True, exist_ok=True)
        archivo = self._carpeta_alertas / \
            f"alertas_{momento.strftime('%Y%m%d_%H%M%S')}.csv"
        with archivo.open("w", newline="", encoding="utf-8-sig") as salida:
            escritor = csv.writer(salida, lineterminator="\n")
            escritor.writerow(self.ENCABEZADOS_ALERTAS)
            for alerta in sorted(alertas,
                                 key=lambda c: -abs(c.diferencia_porcentual)):
                escritor.writerow([
                    alerta.fecha_extraccion.strftime("%Y-%m-%d %H:%M:%S"),
                    alerta.libro.id, alerta.libro.isbn, alerta.libro.titulo,
                    f"{alerta.precio_interno:.2f}",
                    f"{alerta.precio_web.precio:.2f}",
                    f"{alerta.diferencia:.2f}",
                    f"{alerta.diferencia_porcentual:.2f}",
                    f"{umbral_porcentual:.2f}", alerta.tipo_alerta,
                    alerta.precio_web.titulo, alerta.precio_web.url])
        return ResultadoAlertas(archivo, umbral_porcentual, comparaciones,
                                alertas)

    @auditar("SCRAPING CON ALERTAS")
    def scraping_con_alertas(self, umbral_porcentual: float,
                             libros_ids: Optional[Sequence[int]] = None,
                             tipo_dolar_id: Optional[int] = None
                             ) -> Tuple[ResultadoScraping, ResultadoAlertas]:
        """Ejecuta el scraper y genera el CSV de alertas de esa ejecución.

        El umbral se pide antes de ejecutar el scraper para que cada
        ejecución deje su propio archivo de alertas.
        """
        if umbral_porcentual < 0:
            raise ValueError("La diferencia de precios no puede ser negativa.")
        resultado = self._scraping.ejecutar_scraping(libros_ids)
        alertas = self.generar_alertas_csv(umbral_porcentual, tipo_dolar_id,
                                           resultado.fecha_extraccion)
        return resultado, alertas

    # -- Ejercicio 05: reporte Excel ------------------------------------------

    @auditar("REPORTE EXCEL")
    def generar_reporte_excel(self, tipo_dolar_id: Optional[int] = None,
                              fecha: Optional[datetime.datetime] = None,
                              carpeta: Path = CARPETA_REPORTES
                              ) -> Tuple[Path, int]:
        """Genera el reporte Excel de precios internos contra la web.

        La hoja "Comparación" tiene, por libro: Libro, Precio interno,
        Precio web, Diferencia y Fecha de extracción (más el porcentaje,
        el ISBN y el enlace al producto). La hoja "Resultados web" lista
        los 10 resultados de cada búsqueda con sus formas de pago, imagen
        y descripción.

        Returns:
            Tuple[Path, int]: Ruta del archivo y cantidad de libros.
        """
        momento = fecha or self._scraping.ultima_fecha_extraccion()
        comparaciones = self.comparar(tipo_dolar_id, momento)
        if not comparaciones:
            raise ValueError("No hay libros encontrados en la web para "
                             "comparar: ejecute el scraping primero.")
        excel = Workbook()
        hoja = excel.active
        hoja.title = "Comparación"
        self._hoja_comparacion(hoja, comparaciones)
        self._hoja_resultados(excel.create_sheet("Resultados web"),
                              self._scraping.resultados(momento))
        carpeta = Path(carpeta)
        carpeta.mkdir(parents=True, exist_ok=True)
        archivo = carpeta / \
            f"reporte_precios_{momento.strftime('%Y%m%d_%H%M%S')}.xlsx"
        excel.save(archivo)
        return archivo, len(comparaciones)

    FORMATO_PESOS = '"$" #,##0.00'
    FORMATO_FECHA_HORA = "dd/mm/yyyy hh:mm"

    @staticmethod
    def _encabezado(hoja: Worksheet, titulos: Sequence[str],
                    anchos: Sequence[int]) -> None:
        hoja.append(list(titulos))
        relleno = PatternFill("solid", fgColor="1F3864")
        for columna, ancho in enumerate(anchos, start=1):
            celda = hoja.cell(row=1, column=columna)
            celda.font = Font(bold=True, color="FFFFFF")
            celda.fill = relleno
            celda.alignment = Alignment(horizontal="center",
                                        vertical="center", wrap_text=True)
            hoja.column_dimensions[get_column_letter(columna)].width = ancho
        hoja.freeze_panes = "A2"

    def _hoja_comparacion(self, hoja: Worksheet,
                          comparaciones: Sequence[ComparacionPrecio]) -> None:
        self._encabezado(
            hoja, ("Libro", "Precio interno", "Precio web", "Diferencia",
                   "Fecha de extracción", "Diferencia %", "ISBN",
                   "Título en la web", "URL"),
            (38, 16, 16, 16, 20, 14, 16, 45, 60))
        for c in comparaciones:
            hoja.append([c.libro.titulo, c.precio_interno, c.precio_web.precio,
                         c.diferencia, c.fecha_extraccion,
                         c.diferencia_porcentual / 100, c.libro.isbn,
                         c.precio_web.titulo, c.precio_web.url])
            fila = hoja.max_row
            for columna in (2, 3, 4):
                hoja.cell(row=fila, column=columna).number_format = \
                    self.FORMATO_PESOS
            hoja.cell(row=fila, column=5).number_format = \
                self.FORMATO_FECHA_HORA
            hoja.cell(row=fila, column=6).number_format = "0.00%"
            color = "C00000" if c.diferencia < 0 else "00804A"
            for columna in (4, 6):
                hoja.cell(row=fila, column=columna).font = Font(color=color)
        hoja.auto_filter.ref = hoja.dimensions

    def _hoja_resultados(self, hoja: Worksheet,
                         resultados: Sequence[PrecioWeb]) -> None:
        self._encabezado(
            hoja, ("Libro ID", "Libro", "Posición", "Título en la web",
                   "Precio web", "Precio U$s", "Coincidencia",
                   "Seleccionado", "Formas de pago y precio", "URL imagen",
                   "Descripción", "URL", "Fecha de extracción"),
            (9, 30, 9, 40, 14, 11, 12, 12, 60, 45, 70, 50, 20))
        nombres: Dict[int, str] = {}
        for r in resultados:
            if r.libro_id not in nombres:
                try:
                    nombres[r.libro_id] = self._libros.obtener(
                        r.libro_id).titulo
                except ValueError:
                    nombres[r.libro_id] = f"? ({r.libro_id})"
            hoja.append([r.libro_id, nombres[r.libro_id], r.posicion,
                         r.titulo, r.precio, r.precio_usd,
                         r.coincidencia, "Sí" if r.seleccionado else "",
                         "\n".join(r.formas_pago), r.url_imagen,
                         r.descripcion, r.url, r.fecha_extraccion])
            fila = hoja.max_row
            hoja.cell(row=fila, column=5).number_format = self.FORMATO_PESOS
            hoja.cell(row=fila, column=6).number_format = '"U$s" #,##0.00'
            hoja.cell(row=fila, column=13).number_format = \
                self.FORMATO_FECHA_HORA
            for columna in (9, 11):
                hoja.cell(row=fila, column=columna).alignment = Alignment(
                    wrap_text=True, vertical="top")
        hoja.auto_filter.ref = hoja.dimensions


# ---------------------------------------------------------------------------
# Sprint 3 - Ejercicio 06: consulta del historial de auditoría
# ---------------------------------------------------------------------------

class ServicioAuditoria(ServicioBase[Auditoria]):
    """CRUD y consulta del historial de auditoría."""

    nombre_entidad = "un registro de auditoría"
    entidad_auditoria = "Auditoria"

    def __init__(self, repositorio: RepositorioAuditoria) -> None:
        super().__init__(repositorio)
        self._repo_auditoria = repositorio

    @auditar("CONSULTA HISTORIAL")
    def historial(self, cantidad: int = 30,
                  accion: Optional[str] = None) -> List[Auditoria]:
        """Últimas operaciones registradas (las más recientes primero).

        Args:
            cantidad (int): Cantidad máxima de registros a devolver.
            accion (Optional[str]): Filtra por texto en la acción
                (por ejemplo "ALTA", "Libro", "SCRAPING").
        """
        if cantidad <= 0:
            raise ValueError("La cantidad debe ser mayor que cero.")
        return self._repo_auditoria.leer_ultimas(cantidad, accion)


# ---------------------------------------------------------------------------
# Contenedor con todos los servicios del sistema
# ---------------------------------------------------------------------------

@dataclass
class Servicios:
    """Agrupa los servicios para pasarlos a la interfaz de usuario."""

    generos: ServicioGenero
    editoriales: ServicioEditorial
    monedas: ServicioMoneda
    tipos_cotizacion: ServicioTipoCotizacion
    libros: ServicioLibro
    precios: ServicioPrecio
    stock: ServicioStock
    cotizaciones: ServicioCotizacionDolar
    cotizador: ServicioCotizador
    scraping: ServicioScraping
    comparador: ServicioComparadorPrecios
    auditoria: ServicioAuditoria

    @classmethod
    def desde_repositorios(cls, repos: Repositorios) -> "Servicios":
        """Construye todos los servicios inyectando los repositorios.

        También deja configurado el registro de auditoría, de modo que
        todas las operaciones de los servicios queden auditadas.
        """
        RegistroAuditoria.configurar(repos.auditoria)
        libros = ServicioLibro(repos.libros, repos.generos,
                               repos.editoriales, repos.precios, repos.stock)
        precios = ServicioPrecio(repos.precios, repos.libros, repos.monedas)
        monedas = ServicioMoneda(repos.monedas, repos.precios,
                                 repos.tipos_cotizacion)
        tipos = ServicioTipoCotizacion(repos.tipos_cotizacion,
                                       repos.cotizaciones, repos.monedas)
        stock = ServicioStock(repos.stock, repos.libros)
        cotizaciones = ServicioCotizacionDolar(
            repos.cotizaciones, repos.tipos_cotizacion, repos.monedas)
        cotizador = ServicioCotizador(libros, precios, monedas, stock, tipos,
                                      cotizaciones)
        scraping = ServicioScraping(repos.precios_web, repos.libros)
        return cls(
            generos=ServicioGenero(repos.generos, repos.libros),
            editoriales=ServicioEditorial(repos.editoriales, repos.libros),
            monedas=monedas,
            tipos_cotizacion=tipos,
            libros=libros,
            precios=precios,
            stock=stock,
            cotizaciones=cotizaciones,
            cotizador=cotizador,
            scraping=scraping,
            comparador=ServicioComparadorPrecios(libros, monedas, tipos,
                                                 cotizador, scraping),
            auditoria=ServicioAuditoria(repos.auditoria),
        )
